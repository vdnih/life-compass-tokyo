"""ライフコンパス東京 — First Stage プレゼン資料のビルドスクリプト。

このファイルがスライドの唯一のソース。`docs/hackathon/PRESENTATION_OUTLINE.md`
の文言と逐語一致させること。文言を直すときはこのファイルと骨子の両方を直す。

実行方法（リポジトリの依存には加えない。都度 uv で python-pptx を渡して実行する）:

    uv run --with python-pptx docs/hackathon/slides/build_slides.py

出力: docs/hackathon/slides/lifecompass-slides.pptx
"""

from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

HERE = Path(__file__).parent
ASSETS = HERE / "assets"
OUTPUT = HERE / "lifecompass-slides.pptx"

# 16:9, 骨子と同じ比率
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x20, 0x24, 0x2B)
INK_SOFT = RGBColor(0x56, 0x5B, 0x62)
NAVY = RGBColor(0x1E, 0x3A, 0x5F)
NAVY_DEEP = RGBColor(0x14, 0x28, 0x42)
LINE = RGBColor(0xD9, 0xD0, 0xBC)
BLOOM = RGBColor(0xC2, 0x70, 0x8C)

FONT_JP = "Hiragino Sans"  # macOS 前提（発表者の環境）。Windows で開く場合はフォント未搭載のため
                            # OS 側のフォント置換に委ねられる（Yu Gothic 等に自動代替される）

# 事務局「提出資料について」が資料内容の必須項目として求めるチーム紹介。
# 提出フォーム 1-1／1-8 と同じ値を入れる（SUBMISSION_DRAFT.md「1. チーム情報」が正）。
TEAM_NAME = "浦安ライフ"
TEAM_MEMBERS = "あかり（プロダクトオーナー・マーケター）／りょう（エンジニア）"
TEAM_ONELINER = "浦安在住・東京で働く夫婦。「女性がもっとイキイキできる社会に」という思いから作りました。"


def set_font(run, size, color=INK, bold=False, name=FONT_JP):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    # 日本語(East Asian)ランにも明示しないと、macOS でも英数と和文でフォントが分かれて描画されうる
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def add_textbox(slide, left, top, width, height, text, size, color=INK, bold=False,
                 align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.25):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        run = p.add_run()
        run.text = line
        set_font(run, size, color=color, bold=bold)
    return box


def add_rect(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.shadow.inherit = False
    if fill_color is not None:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if line_color is not None:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(0.75)
    else:
        shape.line.fill.background()
    return shape


def add_footnote(slide, text):
    add_textbox(
        slide, Inches(0.6), Inches(6.95), Inches(12.15), Inches(0.5),
        text, 10, color=INK_SOFT, line_spacing=1.15,
    )


def add_slide_number(slide, n):
    add_textbox(
        slide, Inches(12.55), Inches(0.25), Inches(0.6), Inches(0.4),
        f"{n:02d}", 12, color=INK_SOFT, align=PP_ALIGN.RIGHT,
    )


def add_picture_framed(slide, path, left, top, width, height):
    """画像を枠線付きで配置する（実画面のスクリーンショットであることを示す）。

    アセットが見当たらない場合は、空白のまま埋め込まれて気づかれないより、
    ビルドを止めて気づけるほうがよい。
    """
    if not path.exists():
        raise FileNotFoundError(f"スライド用アセットが見つかりません: {path}")
    pic = slide.shapes.add_picture(str(path), left, top, width=width, height=height)
    pic.line.color.rgb = LINE
    pic.line.width = Pt(1)
    return pic


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # 白紙レイアウト
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    return slide


def set_notes(slide, text):
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    tf.text = text


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ------------------------------------------------------------------
    # 1. 表紙｜ライフコンパス東京
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_rect(s, 0, 0, Inches(0.18), SLIDE_H, fill_color=NAVY)
    add_textbox(s, Inches(0.7), Inches(0.9), Inches(6.0), Inches(1.2),
                "ライフコンパス東京", 44, color=NAVY_DEEP, bold=True)
    add_textbox(s, Inches(0.7), Inches(1.95), Inches(6.0), Inches(0.8),
                "自分の人生を、自分で描き直す", 22, color=NAVY, bold=True)
    add_textbox(s, Inches(0.7), Inches(2.75), Inches(6.0), Inches(0.5),
                "LIFECOMPASS TOKYO", 12, color=INK_SOFT)
    add_picture_framed(s, ASSETS / "timeline-overview.jpg",
                        Inches(7.0), Inches(1.4), width=Inches(5.8), height=Inches(3.31))

    # 事務局「提出資料について」で資料内容の必須項目として求められているチーム紹介。
    # 本編2分を圧迫しないよう、独立スライドではなく表紙に置く。
    add_rect(s, Inches(0.7), Inches(3.55), Inches(5.9), Inches(1.65),
             fill_color=RGBColor(0xF2, 0xF4, 0xF7))
    add_textbox(s, Inches(0.95), Inches(3.7), Inches(5.4), Inches(0.4),
                f"チーム {TEAM_NAME}（2名）", 14, color=NAVY, bold=True)
    add_textbox(s, Inches(0.95), Inches(4.1), Inches(5.4), Inches(1.0),
                f"{TEAM_MEMBERS}\n{TEAM_ONELINER}",
                12, color=INK, line_spacing=1.35)

    add_footnote(s, "都知事杯オープンデータ・ハッカソン2026 First Stage")
    add_slide_number(s, 1)
    set_notes(s, "[話者: あかり / 0:00-0:13]\n"
                 "浦安ライフのあかりです。エンジニアの夫、りょうと作りました。"
                 "女性がもっとイキイキできる社会に。その思いで作った、ライフコンパス東京です。")

    # ------------------------------------------------------------------
    # 2. 課題①｜家庭もキャリアも、全力でがんばりたい／でも周りはこう言う
    #
    # 見出しは必ず「彼女」を主語に置く。ここを「思い込み」側の主語にすると、
    # 全体メッセージの起点（家庭もキャリアも全力でがんばりたい女性がいる）が
    # 資料から消える。
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(11.5), Inches(0.8),
                "家庭もキャリアも、全力でがんばりたい", 34, color=NAVY_DEEP, bold=True)
    add_textbox(s, Inches(0.7), Inches(1.2), Inches(11.5), Inches(0.5),
                "——でも、周りは「女性は家庭の方が向いている」と言う",
                20, color=BLOOM, bold=True)

    add_textbox(s, Inches(0.7), Inches(1.85), Inches(11.5), Inches(0.35),
                "「育児は◯◯の方が向いている」と思う人の割合", 14, color=INK_SOFT)

    # 左: 女性の方が向いている
    add_textbox(s, Inches(0.9), Inches(2.3), Inches(5.4), Inches(0.5),
                "育児は女性の方が向いている", 20, color=BLOOM, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(s, Inches(0.9), Inches(2.85), Inches(5.4), Inches(1.2),
                "68%", 58, color=BLOOM, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(s, Inches(6.35), Inches(3.15), Inches(0.7), Inches(0.6),
                "↔", 32, color=INK_SOFT, align=PP_ALIGN.CENTER)

    # 右: 男性の方が向いている
    add_textbox(s, Inches(7.05), Inches(2.3), Inches(5.4), Inches(0.5),
                "育児は男性の方が向いている", 20, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(s, Inches(7.05), Inches(2.85), Inches(5.4), Inches(1.2),
                "7%", 58, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(s, Inches(0.7), Inches(4.35), Inches(11.5), Inches(0.6),
                "望んでいないのではなく、諦めざるを得ないのではないか。",
                24, color=NAVY_DEEP, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(s, Inches(0.7), Inches(5.05), Inches(11.5), Inches(1.1),
                "企業が挙げる課題1位は「管理職を希望する女性が少ない」46.4%。\n"
                "でも本人が挙げる1位は「家庭責任が重いイメージ」69.4%。",
                17, color=INK, line_spacing=1.4, align=PP_ALIGN.CENTER)

    add_footnote(
        s,
        "出典: 東京都「男女平等参画に関する世論調査」（令和7年8月調査, n=1,615。"
        "「性別で向いている仕事・向いていない仕事がある」は83%）／"
        "東京都「男女雇用平等参画状況調査」（令和7年度、事業所 n=347・従業員女性 n=556）",
    )
    add_slide_number(s, 2)
    set_notes(s, "[話者: あかり / 0:13-0:39]\n"
                 "家庭もキャリアも、全力でがんばりたい。そういう女性が、身近にいます。"
                 "でも、育児は女性が向いていると考える人が68パーセント、男性は7パーセント。"
                 "企業は「女性が望んでいない」と見る。でも本人の悩みは「家庭責任が重いイメージ」。"
                 "望んでいないのではなく、諦めているのではないでしょうか。"
                 "\n\n※サブ見出しは読み上げない（スライドの読み上げにならないようにする）。")

    # ------------------------------------------------------------------
    # 3. 課題②｜がんばりたくても、使える制度を知らない
    #
    # スライド2と同じく主語は「彼女」。2枚で一続きの文として読ませる。
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.45), Inches(11.5), Inches(0.8),
                "がんばりたくても、使える制度を知らない", 34, color=NAVY_DEEP, bold=True)
    add_textbox(s, Inches(0.7), Inches(1.2), Inches(11.5), Inches(0.4),
                "自分の職場に制度があるか「わからない」", 16, color=INK_SOFT)

    bars = [
        ("産前産後休暇", 6.0, NAVY),
        ("通院休暇制度", 33.9, NAVY),
        ("妊娠障害休暇", 42.8, NAVY),
        ("出産障害休暇", 47.3, BLOOM),
    ]
    bar_left = Inches(3.1)
    bar_max = Inches(6.2)  # 50%を満尺として換算
    row_top = 1.85
    row_height = 0.85
    for label, pct, color in bars:
        top = Inches(row_top)
        add_textbox(s, Inches(0.7), top + Inches(0.08), Inches(2.2), Inches(0.5),
                    label, 16, color=INK, bold=(color == BLOOM))
        bar_width = Inches(bar_max.inches * pct / 50)
        add_rect(s, bar_left, top, bar_width, Inches(0.5), fill_color=color)
        add_textbox(s, bar_left + bar_width + Inches(0.15), top + Inches(0.03), Inches(1.5),
                    Inches(0.5), f"{pct:.1f}%", 18, color=color, bold=True)
        row_top += row_height

    add_textbox(s, Inches(0.7), Inches(5.55), Inches(11.5), Inches(0.6),
                "知らない制度は、選択肢に入らない。",
                24, color=NAVY_DEEP, bold=True, align=PP_ALIGN.CENTER)

    add_footnote(
        s,
        "出典: 東京都「男女雇用平等参画状況調査」（令和7年度）従業員調査（n=1,076）。"
        "母性保護に関する制度ごとの「わからない」の割合",
    )
    add_slide_number(s, 3)
    set_notes(s, "[話者: あかり / 0:39-0:56]\n"
                 "そして、使える制度を知りません。産前産後休暇を知らない人は6パーセント。"
                 "でも出産障害休暇は、47パーセントが「職場にあるか分からない」。"
                 "知らない制度は、選択肢に入りません。")

    # ------------------------------------------------------------------
    # 4. 解決策｜何が選べるかを知って、自分で決める
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.9),
                "何が選べるかを知って、自分で決める", 34, color=NAVY_DEEP, bold=True)

    add_picture_framed(s, ASSETS / "event-detail-childbirth.jpg",
                        Inches(0.7), Inches(1.5), width=Inches(7.4), height=Inches(4.23))

    labels = ["決めるのは自分", "同じ時間軸", "出典付きの事実"]
    for i, label in enumerate(labels):
        top = Inches(1.6 + i * 1.15)
        add_rect(s, Inches(8.5), top, Inches(4.0), Inches(0.85), fill_color=RGBColor(0xF2, 0xF4, 0xF7))
        add_textbox(s, Inches(8.7), top + Inches(0.18), Inches(3.6), Inches(0.5),
                    f"0{i+1}  {label}", 18, color=NAVY, bold=True)

    add_textbox(s, Inches(8.5), Inches(5.3), Inches(4.0), Inches(1.5),
                "「国の制度」「東京都の制度」を実施主体ごとに明示し、\n"
                "公式ページへのリンクを常に添える。",
                14, color=INK_SOFT, line_spacing=1.3)

    add_footnote(
        s,
        "出典: 育児・介護休業法／労働基準法／母子保健法（厚生労働省・こども家庭庁）、"
        "東京都不妊検査等助成事業（東京都福祉局）。平均値・他人との比較は表示しない。",
    )
    add_slide_number(s, 4)
    set_notes(s, "[話者: りょう / 0:56-1:14 ※ここから交代]\n"
                 "ここからは実装です。都の子育て支援制度は7,812件。読み切れる量ではありません。"
                 "そこで、育休は子が2歳まで、といった国と都の制度上の上限だけを、"
                 "実施主体と出典リンク付きで表示しました。")

    # ------------------------------------------------------------------
    # 5. デモ｜何度でも立て直せる
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.9),
                "何度でも立て直せる", 36, color=NAVY_DEEP, bold=True)

    add_picture_framed(s, ASSETS / "timeline-overview.jpg",
                        Inches(0.7), Inches(1.5), width=Inches(4.4), height=Inches(3.25))
    add_picture_framed(s, ASSETS / "timeline-mobile.jpg",
                        Inches(5.3), Inches(1.5), width=Inches(1.5), height=Inches(3.25))
    add_picture_framed(s, ASSETS / "event-detail-childbirth.jpg",
                        Inches(7.0), Inches(1.5), width=Inches(5.3), height=Inches(3.25))

    add_textbox(s, Inches(0.7), Inches(5.0), Inches(11.5), Inches(1.3),
                "結婚・出産・転職などの予定を自由に置き、必要な制度情報をすぐに確認できます。\n"
                "一度決めて終わりではなく、状況が変わるたびに何度でも軽やかに描き直せます。",
                20, color=INK, line_spacing=1.4)

    add_textbox(s, Inches(0.7), Inches(6.55), Inches(6.0), Inches(0.4),
                "公開デモ: my-career-app-559fd.web.app", 14, color=NAVY, bold=True)
    add_footnote(
        s,
        "希望される支援も「新しい休業」より「時間の組み替え」（有給53.7%・フレックス49.4%、"
        "東京都「男女雇用平等参画状況調査」令和7年度）",
    )
    add_slide_number(s, 5)
    set_notes(s, "[話者: りょう / 1:14-1:31]\n"
                 "結婚、出産、転職の予定を、同じ時間軸に自由に置けます。"
                 "ひとつ動かすと、つながるイベントも一緒に動く。"
                 "AIも、確認済みの制度情報だけを根拠に答えます。何度でも描き直せます。"
                 "\n\n[登壇メモ] 収録ではライブデモ（アプリ画面のリアルタイム操作）は一切できない"
                 "（事務局「提出資料について」）。動く様子を見せる場合は、事前に収録した無音動画を"
                 "このスライドに埋め込み、画面共有オプションの2項目にチェックを入れて共有する。"
                 "対話型AI（Gemini Function Calling）の応答画面は未サインインでは取得できないため"
                 "（送信にGoogleサインインが必要）、サインイン状態で録画したものを使う。")

    # ------------------------------------------------------------------
    # 6. 広がり｜一つの画面を、家庭で。男性が体験する
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.9),
                "一つの画面を、家庭で。男性が体験する", 32, color=NAVY_DEEP, bold=True)

    add_textbox(s, Inches(0.7), Inches(1.6), Inches(11.5), Inches(0.6),
                "同じ画面を、二人で見る", 22, color=INK, bold=True, align=PP_ALIGN.CENTER)

    # 左: 実装済み
    add_rect(s, Inches(1.2), Inches(2.6), Inches(5.0), Inches(2.3), fill_color=RGBColor(0xF2, 0xF4, 0xF7))
    add_textbox(s, Inches(1.5), Inches(2.85), Inches(4.4), Inches(0.5),
                "実装済み・ログイン不要", 14, color=NAVY, bold=True)
    add_textbox(s, Inches(1.5), Inches(3.35), Inches(4.4), Inches(1.3),
                "自分のプランを描く", 26, color=NAVY_DEEP, bold=True)

    # 右: 構想
    add_rect(s, Inches(7.1), Inches(2.6), Inches(5.0), Inches(2.3), line_color=NAVY)
    add_textbox(s, Inches(7.4), Inches(2.85), Inches(4.4), Inches(0.5),
                "構想中", 14, color=BLOOM, bold=True)
    add_textbox(s, Inches(7.4), Inches(3.25), Inches(4.4), Inches(1.5),
                "パートナーと一緒に見る\n男性が自分ごととして体験する",
                20, color=NAVY_DEEP, bold=True, line_spacing=1.3)

    add_footnote(
        s,
        "出典: 東京都「男女雇用平等参画状況調査」（令和7年度）。母性保護に関する制度8項目"
        "すべてで、男性の「わからない」が女性より高い（出産障害休暇: 男性52.6%／女性42.4%）。",
    )
    add_slide_number(s, 6)
    set_notes(s, "[話者: りょう / 1:31-1:46]\n"
                 "この画面はログイン不要で、家庭で一緒に見られます。"
                 "制度を「分からない」と答えた割合は、8項目すべてで男性の方が高い。"
                 "男性が体験する場にも広げたいと考えています。"
                 "\n\n[質疑用] 男性の育休取得率は61.2%まで上がったが、取得期間は「1か月〜3か月未満」"
                 "38.6%が最多で、女性の「6か月〜1年未満」30.2%とは差がある。")

    # ------------------------------------------------------------------
    # 7. 締め｜このデータが、次の施策を動かす
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.9),
                "このデータが、次の施策を動かす", 34, color=NAVY_DEEP, bold=True)

    steps = ["東京都・国の\nオープンデータ", "一人ひとりの\nライフプラン", "次の施策へ"]
    step_w = Inches(3.4)
    gap = Inches(0.5)
    start_left = Inches(0.9)
    for i, step in enumerate(steps):
        left = start_left + i * (step_w + gap)
        add_rect(s, left, Inches(1.9), step_w, Inches(1.5), fill_color=RGBColor(0xF2, 0xF4, 0xF7))
        add_textbox(s, left, Inches(2.15), step_w, Inches(1.0), step, 18, color=NAVY_DEEP,
                    bold=True, align=PP_ALIGN.CENTER, line_spacing=1.2)
        if i < len(steps) - 1:
            arrow_left = left + step_w
            add_textbox(s, arrow_left, Inches(2.35), gap, Inches(0.6), "→", 24,
                        color=INK_SOFT, align=PP_ALIGN.CENTER)

    add_textbox(s, Inches(0.7), Inches(4.2), Inches(11.5), Inches(1.0),
                "知ったうえで、自分で決める", 32, color=NAVY, bold=True, align=PP_ALIGN.CENTER)

    add_textbox(s, Inches(0.7), Inches(5.4), Inches(11.5), Inches(0.5),
                "公開デモ: my-career-app-559fd.web.app", 16, color=NAVY_DEEP, bold=True,
                align=PP_ALIGN.CENTER)

    add_footnote(
        s,
        "制度情報の出典は各スライド記載のとおり。行政に力を入れてほしいこと1位「家事・育児や"
        "介護中の人への家庭と仕事の両立支援」64%（東京都「男女平等参画に関する世論調査」令和7年8月調査）",
    )
    add_slide_number(s, 7)
    set_notes(s, "[話者: あかり / 1:46-2:00 ※ここで交代して締める]\n"
                 "誰がどこで諦めたかのデータは、都にも国にもありません。"
                 "ここで積み重なるプランが、次の施策につながることを願っています。"
                 "決めるのは、いつも本人です。")

    # ------------------------------------------------------------------
    # 8. 参考（付録）｜利用データと使い方
    #
    # 提出フォーム 4-1「利用データ一覧」はデータ名とURLのセットしか登録できず、
    # 提供元・ライセンス・使われ方までは書けない。審査5軸の筆頭「データ活用」に
    # 応えるため、その説明をこの付録ページに置く。本編2分では話さない。
    # ------------------------------------------------------------------
    s = new_slide(prs)
    add_textbox(s, Inches(0.7), Inches(0.5), Inches(11.5), Inches(0.9),
                "参考｜利用データと使い方", 30, color=NAVY_DEEP, bold=True)

    # (B) アプリ内に出典付きで表示するデータ
    add_rect(s, Inches(0.7), Inches(1.5), Inches(11.9), Inches(2.05),
             fill_color=RGBColor(0xF2, 0xF4, 0xF7))
    add_textbox(s, Inches(0.95), Inches(1.65), Inches(11.4), Inches(0.4),
                "(B) アプリ内に出典付きで表示するデータ", 16, color=NAVY, bold=True)
    add_textbox(
        s, Inches(0.95), Inches(2.1), Inches(11.15), Inches(1.35),
        "・東京都不妊検査等助成事業／東京都福祉局／公的な制度情報\n"
        "・育児・介護休業法、労働基準法、母子保健法／厚生労働省・こども家庭庁／公的な法令情報",
        14, color=INK, line_spacing=1.4,
    )

    # (A) 課題の根拠として資料で引用するデータ（利用者の画面には出さない）
    add_rect(s, Inches(0.7), Inches(3.75), Inches(11.9), Inches(2.55), line_color=LINE)
    add_textbox(s, Inches(0.95), Inches(3.9), Inches(11.4), Inches(0.4),
                "(A) 課題の根拠として資料で引用するデータ（利用者の画面には出さない）", 16,
                color=NAVY, bold=True)
    add_textbox(
        s, Inches(0.95), Inches(4.35), Inches(11.15), Inches(1.85),
        "・男女雇用平等参画状況調査（令和7年度）／東京都産業労働局／CC BY 4.0\n"
        "・男女平等参画に関する世論調査（令和7年8月調査）／東京都政策企画局\n"
        "・子育て支援制度レジストリ（東京デジタル2030ビジョン こどもDX）／東京都デジタルサービス局／\n"
        "  CC BY 4.0（7,812件の根拠として引用。更新終了のため制度上限の一次情報は東京都福祉局の\n"
        "  公式ページを採用）",
        14, color=INK, line_spacing=1.4,
    )

    add_textbox(s, Inches(0.7), Inches(6.45), Inches(11.9), Inches(0.5),
                "平均値・中央値は画面に出さず、制度上の上限だけを実施主体・出典付きで示す（PDR-008）",
                14, color=NAVY, bold=True)

    add_footnote(s, "各データの出典URL・ライセンス表記の詳細は提出フォーム 4-1 に記載。")
    add_slide_number(s, 8)
    set_notes(s, "[付録：本編2分では話さない。質疑で聞かれたときのための補足]"
                 "アプリ画面に表示する制度上限は、東京都福祉局の不妊検査等助成事業と、"
                 "育児・介護休業法など国の法令から出典付きで取り込んでいます。"
                 "課題の根拠として資料・プレゼンで引用した東京都の調査データは、"
                 "利用者の画面には表示していません。")

    prs.save(str(OUTPUT))
    print(f"wrote {OUTPUT}")


if __name__ == "__main__":
    build()
